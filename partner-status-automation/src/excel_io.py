"""
Excel I/O: openpyxl-based read/write, hidden column/row detection.
All cell-level work uses openpyxl directly to preserve formatting.
pandas is used only for display DataFrames.
"""
from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.utils import col_letter_to_index, safe_str


# ---------------------------------------------------------------------------
# Sheet utilities
# ---------------------------------------------------------------------------

def get_sheet_names(file_path: str) -> List[str]:
    """Return the list of sheet names in *file_path*."""
    wb = load_workbook(file_path, read_only=True, data_only=True)
    names = wb.sheetnames
    wb.close()
    return names


# ---------------------------------------------------------------------------
# Hidden column / row detection
# ---------------------------------------------------------------------------

def detect_hidden_columns(ws: Worksheet) -> List[str]:
    """Return sorted list of hidden column letters in *ws*."""
    hidden: List[str] = []
    for col_letter, col_dim in ws.column_dimensions.items():
        if col_dim.hidden:
            hidden.append(col_letter)
    # Sort by column index
    hidden.sort(key=lambda c: col_letter_to_index(c))
    return hidden


def detect_hidden_rows(ws: Worksheet) -> List[int]:
    """Return sorted list of hidden row numbers in *ws*."""
    hidden: List[int] = []
    for row_num, row_dim in ws.row_dimensions.items():
        if row_dim.hidden:
            hidden.append(row_num)
    hidden.sort()
    return hidden


def get_hidden_info(file_path: str, sheet_name: str) -> Tuple[List[str], List[int]]:
    """
    Return (hidden_columns, hidden_rows) for the given sheet.
    Reads in normal mode so column_dimensions is populated.
    """
    wb = load_workbook(file_path, data_only=True)
    ws = wb[sheet_name]
    hidden_cols = detect_hidden_columns(ws)
    hidden_rows = detect_hidden_rows(ws)
    wb.close()
    return hidden_cols, hidden_rows


# ---------------------------------------------------------------------------
# Data extraction (openpyxl)
# ---------------------------------------------------------------------------

def _cell_value(cell) -> Any:
    """Return raw cell value; None for empty cells."""
    return cell.value


def read_sheet_data(
    file_path: str,
    sheet_name: str,
    key_col: str,
    value_cols: Optional[List[str]] = None,
) -> Tuple[Dict[str, Dict[str, Any]], List[int]]:
    """
    Read rows from *sheet_name* where the key column has a non-empty value.

    Returns
    -------
    data : dict  key_value → {col_letter: raw_cell_value, ...}
    row_numbers : list of row numbers corresponding to each key (same order)

    Hidden columns/rows are included — we read everything openpyxl sees.
    """
    wb = load_workbook(file_path, data_only=True)
    ws = wb[sheet_name]

    key_idx = col_letter_to_index(key_col)
    value_indices: Dict[str, int] = {}
    if value_cols:
        for vc in value_cols:
            value_indices[vc] = col_letter_to_index(vc)

    data: Dict[str, Dict[str, Any]] = {}
    row_numbers: List[int] = []

    for row in ws.iter_rows():
        # Get key cell
        key_cell = None
        for cell in row:
            if cell.column == key_idx:
                key_cell = cell
                break
        if key_cell is None:
            continue
        key_val = safe_str(_cell_value(key_cell))
        if not key_val:
            continue

        row_data: Dict[str, Any] = {}
        if value_cols:
            for col_letter, col_idx in value_indices.items():
                for cell in row:
                    if cell.column == col_idx:
                        row_data[col_letter] = _cell_value(cell)
                        break
                else:
                    row_data[col_letter] = None

        if key_val not in data:
            data[key_val] = row_data
            row_numbers.append(key_cell.row)
        # Duplicates: keep first occurrence (caller handles dup detection)

    wb.close()
    return data, row_numbers


def read_sheet_data_with_duplicates(
    file_path: str,
    sheet_name: str,
    key_col: str,
    value_cols: Optional[List[str]] = None,
) -> Tuple[Dict[str, List[Dict[str, Any]]], List[int]]:
    """
    Like read_sheet_data but stores ALL occurrences of each key.
    Returns {key: [row_data, ...]} so callers can detect duplicates.
    """
    wb = load_workbook(file_path, data_only=True)
    ws = wb[sheet_name]

    key_idx = col_letter_to_index(key_col)
    value_indices: Dict[str, int] = {}
    if value_cols:
        for vc in value_cols:
            value_indices[vc] = col_letter_to_index(vc)

    data: Dict[str, List[Dict[str, Any]]] = {}
    row_numbers: List[int] = []

    for row in ws.iter_rows():
        key_cell = None
        for cell in row:
            if cell.column == key_idx:
                key_cell = cell
                break
        if key_cell is None:
            continue
        key_val = safe_str(_cell_value(key_cell))
        if not key_val:
            continue

        row_data: Dict[str, Any] = {"_row": key_cell.row}
        if value_cols:
            for col_letter, col_idx in value_indices.items():
                for cell in row:
                    if cell.column == col_idx:
                        row_data[col_letter] = _cell_value(cell)
                        break
                else:
                    row_data[col_letter] = None

        data.setdefault(key_val, []).append(row_data)
        row_numbers.append(key_cell.row)

    wb.close()
    return data, row_numbers


def read_keys_only(file_path: str, sheet_name: str, key_col: str) -> List[str]:
    """Return the list of non-empty key values in key_col order."""
    wb = load_workbook(file_path, data_only=True)
    ws = wb[sheet_name]
    key_idx = col_letter_to_index(key_col)
    keys: List[str] = []
    for row in ws.iter_rows(min_col=key_idx, max_col=key_idx):
        val = safe_str(row[0].value)
        if val:
            keys.append(val)
    wb.close()
    return keys


# ---------------------------------------------------------------------------
# Writing changes back to the original workbook
# ---------------------------------------------------------------------------

def save_with_changes(
    source_path: str,
    output_path: str,
    sheet_name: str,
    changes: List[Dict],
) -> None:
    """
    Copy *source_path* → *output_path*, then apply *changes* to the worksheet.

    Each element of *changes* is:
        {"row": int, "col_letter": str, "value": any}

    Preserves all formatting, filters, column widths, hidden states, merges.
    """
    # Make a copy so we never overwrite the original
    shutil.copy2(source_path, output_path)

    wb = load_workbook(output_path)
    ws = wb[sheet_name]

    for change in changes:
        row = change["row"]
        col_idx = col_letter_to_index(change["col_letter"])
        value = change["value"]
        cell = ws.cell(row=row, column=col_idx)
        cell.value = value

    wb.save(output_path)
    wb.close()


# ---------------------------------------------------------------------------
# Read all rows with row numbers for final file (used by matcher)
# ---------------------------------------------------------------------------

def read_final_file_rows(
    file_path: str,
    sheet_name: str,
    key_col: str,
) -> List[Tuple[int, str]]:
    """
    Return list of (row_number, key_value) for all rows with a non-empty key.
    Row numbers are 1-based (openpyxl convention).
    """
    wb = load_workbook(file_path, data_only=True)
    ws = wb[sheet_name]
    key_idx = col_letter_to_index(key_col)

    result: List[Tuple[int, str]] = []
    for row in ws.iter_rows():
        key_cell = None
        for cell in row:
            if cell.column == key_idx:
                key_cell = cell
                break
        if key_cell is None:
            continue
        key_val = safe_str(_cell_value(key_cell))
        if key_val:
            result.append((key_cell.row, key_val))

    wb.close()
    return result
